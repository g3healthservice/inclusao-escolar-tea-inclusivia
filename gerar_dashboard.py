#!/usr/bin/env python3
"""Gera o index.html autocontido do painel Inclusi.Via (TEA).

Entradas:
  template.html                          layout + lógica (placeholders __X__)
  dados/tea_municipios_censo2022.json    [codIBGE, UF, nome, 0-4, 5-9, 10-14, 15-19]
  dados/censo_escolar_2025.json          educação especial por município (INEP) — importar_censo_escolar.py
  dados/siconfi_educacao_especial.json   gasto com Educação Especial por ente (Tesouro) — coleta_siconfi.py
  dados/escolas_resumo.json + escolas/   mapa de escolas (microdados INEP) — importar_escolas.py; escolas/<UF>.json é servido junto do site
  vendor/*.js                            Chart.js, jsPDF, jsPDF-AutoTable
  mail_endpoint.txt / mail_token.txt     (opcionais) Apps Script de envio de e-mail

Atualizar a base: python3 gerar_dashboard.py --xlsx ~/Downloads/autismo_municipios_censo2022.xlsx
"""
import argparse, json, os, sys
from datetime import datetime

BASE = os.path.dirname(os.path.abspath(__file__))
SITE_URL = "https://g3healthservice.github.io/inclusao-escolar-tea-inclusivia/"


def ler(p):
    with open(os.path.join(BASE, p), encoding="utf-8") as f:
        return f.read()


def ler_opcional(p):
    fp = os.path.join(BASE, p)
    return open(fp, encoding="utf-8").read().strip() if os.path.exists(fp) else ""


def importar_xlsx(caminho):
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Código IBGE")
    data = [[str(r[0]), r[1], r[2], int(r[4] or 0), int(r[5] or 0), int(r[6] or 0), int(r[7] or 0)]
            for r in rows[hi + 1:] if r[0]]
    with open(os.path.join(BASE, "dados/tea_municipios_censo2022.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"Base importada: {len(data)} municípios, {sum(sum(d[3:]) for d in data):,} pessoas com TEA 0-19")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="reimporta a planilha do Censo antes de gerar")
    ap.add_argument("--out", default="index.html")
    a = ap.parse_args()
    if a.xlsx:
        importar_xlsx(os.path.expanduser(a.xlsx))

    dados = ler("dados/tea_municipios_censo2022.json")
    lib = lambda n: ler("vendor/" + n).replace("</script", "<\\/script")
    endpoint, token = ler_opcional("mail_endpoint.txt"), ler_opcional("mail_token.txt")
    build = datetime.now().strftime("%Y-%m-%d %H:%M")

    docs = ler("docs.js").replace("__XLSX_STATIC__", ler("dados/xlsx_estatico.json"))
    html = ler("template.html").replace("__DOCSJS__", docs)
    # dados e libs por último: não podem ter seus conteúdos reinterpretados como placeholder
    for k, v in {"__SITE_URL__": SITE_URL, "__BUILD__": build,
                 "__MAIL_ENDPOINT__": endpoint, "__MAIL_TOKEN__": token}.items():
        html = html.replace(k, v)
    import base64
    html = html.replace("__BRAINLOGO__", "data:image/png;base64," + base64.b64encode(open(os.path.join(BASE, "brain-logo.png"), "rb").read()).decode())
    html = html.replace("__DATA__", dados)
    html = html.replace("__CENSO__", ler_opcional("dados/censo_escolar_2025.json") or "{}")
    html = html.replace("__SICONFI__", ler_opcional("dados/siconfi_educacao_especial.json") or "{}")
    html = html.replace("__ESCOLAS__", ler_opcional("dados/escolas_resumo.json") or "{}")
    html = html.replace("__CHARTJS__", lib("chart.umd.js"))
    html = html.replace("__JSPDF__", lib("jspdf.umd.min.js"))
    html = html.replace("__AUTOTABLE__", lib("jspdf.plugin.autotable.min.js"))

    out = os.path.join(BASE, a.out)
    with open(out, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"OK {out} ({len(html)/1024:.0f} KB) · build {build} · e-mail: {'CONFIGURADO' if endpoint else 'mailto (sem endpoint)'}")


if __name__ == "__main__":
    sys.exit(main())
